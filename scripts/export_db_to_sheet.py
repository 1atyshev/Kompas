import sqlite3
from pathlib import Path
from typing import Iterable

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "bot.db"
TEMPLATE_PATH = ROOT / "habitbot_template_EN.xlsx"


def clear_range(ws, start_row: int, start_col: int, end_col: int) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.value = None


def write_rows(ws, start_row: int, rows: Iterable[tuple]) -> None:
    row_idx = start_row
    for row in rows:
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx).value = value
        row_idx += 1


def main() -> None:
    if not DB_PATH.exists():
        raise SystemExit(f"DB not found: {DB_PATH}")
    if not TEMPLATE_PATH.exists():
        raise SystemExit(f"Template not found: {TEMPLATE_PATH}")

    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    cur.execute(
        "SELECT id, name, type, unit, xp, target, period, active, streak_mode FROM trackers ORDER BY order_index, id"
    )
    trackers = cur.fetchall()

    cur.execute("SELECT tracker_id, period, target FROM tracker_goals ORDER BY tracker_id, period")
    goals = cur.fetchall()

    cur.execute(
        "SELECT id, tracker_id, date, value, partial FROM tracker_logs ORDER BY date, tracker_id, id"
    )
    logs = cur.fetchall()

    journal_entries = []
    weekly_reports = []
    try:
        cur.execute(
            "SELECT date, text, updated_at, created_at FROM journal_entries ORDER BY date"
        )
        journal_entries = cur.fetchall()
    except sqlite3.Error:
        journal_entries = []
    try:
        cur.execute(
            "SELECT week_start, week_end, analysis_text, created_at FROM journal_weekly_reports ORDER BY week_start"
        )
        weekly_reports = cur.fetchall()
    except sqlite3.Error:
        weekly_reports = []

    conn.close()

    wb = openpyxl.load_workbook(TEMPLATE_PATH)

    # raw_trackers
    ws = wb["raw_trackers"]
    clear_range(ws, 2, 1, 8)
    rows = []
    for t in trackers:
        target_day = t["target"] if t["period"] == "day" and t["target"] else None
        rows.append(
            (
                t["id"],
                t["name"],
                t["type"],
                t["unit"] or "",
                t["xp"],
                target_day,
                bool(t["active"]),
                t["streak_mode"] or "activity",
            )
        )
    write_rows(ws, 2, rows)

    # raw_tracker_goals
    ws = wb["raw_tracker_goals"]
    clear_range(ws, 2, 1, 3)
    rows = [(g["tracker_id"], g["period"], g["target"]) for g in goals]
    write_rows(ws, 2, rows)

    # raw_tracker_logs
    ws = wb["raw_tracker_logs"]
    clear_range(ws, 2, 1, 5)
    rows = [
        (l["id"], l["tracker_id"], l["date"], float(l["value"]), bool(l["partial"]))
        for l in logs
    ]
    write_rows(ws, 2, rows)

    # journal
    sheet_title = "Дневник"
    if sheet_title in wb.sheetnames:
        ws = wb[sheet_title]
    else:
        ws = wb.create_sheet(sheet_title)
    ws.delete_rows(1, ws.max_row)
    ws.append(["type", "date", "week_start", "week_end", "text", "created_at"])
    rows = []
    for row in journal_entries:
        created = row["updated_at"] or row["created_at"] or ""
        rows.append(("daily", row["date"], "", "", row["text"], created))
    for row in weekly_reports:
        created = row["created_at"] or ""
        rows.append(("weekly", "", row["week_start"], row["week_end"], row["analysis_text"], created))
    for r in rows:
        ws.append(list(r))

    wb.save(TEMPLATE_PATH)
    print(
        f"Exported {len(trackers)} trackers, {len(goals)} goals, {len(logs)} logs, "
        f"{len(journal_entries)} journal entries, {len(weekly_reports)} weekly reports"
    )


if __name__ == "__main__":
    main()
